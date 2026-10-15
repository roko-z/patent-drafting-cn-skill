#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""中国发明专利 DOCX 第一轮结构化审查（非法律结论）。

用法：
 python patent_lint.py --claims 权利要求书.docx --spec 说明书.docx --abstract 说明书摘要.docx --out 报告.md
依赖：pip install python-docx
"""
import argparse
import re
from pathlib import Path
from docx import Document

HEADINGS = ("技术领域", "背景技术", "发明内容", "附图说明", "具体实施方式")
NUM = re.compile(r"^\s*(\d{1,3})\s*[.．、]\s*(.+)", re.S)
REF = re.compile(r"权利要求\s*(\d{1,3})")
FIG = re.compile(r"图\s*(\d{1,3})(?!\d)")
PLACEHOLDER = re.compile(r"【[^】]{1,120}】|\b(?:TODO|TBD|FIXME)\b", re.I)
WORDY = re.compile(r"(世界首创|绝对领先|百分之百|保证授权|无任何缺点)")


def contents(doc):
    result = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                result.extend(p.text.strip() for p in c.paragraphs if p.text.strip())
    return result


def line_add(out, severity, loc, msg):
    out.append((severity, loc, msg))


def check_base(path, label, out):
    doc = Document(path)
    for s in doc.sections:
        # 版心仅为自动预警，法规中底边留白涉及页码实际位置。
        dims = (("顶部", s.top_margin.cm, 2.5), ("左侧", s.left_margin.cm, 2.5),
                ("右侧", s.right_margin.cm, 1.5), ("底部", s.bottom_margin.cm, 1.5))
        for name, actual, expected in dims:
            if abs(actual - expected) > 0.11:
                line_add(out, "提示", label, f"{name} Word 页边距 {actual:.2f}cm；内部常用起点为 {expected:.2f}cm。仍应核对现行官方规则。")
        if abs(s.page_width.cm - 21.0) > .1 or abs(s.page_height.cm - 29.7) > .1:
            line_add(out, "重大", label, "不是常用A4纵向纸张（21 × 29.7 cm）。")
    for i, text in enumerate(contents(doc), 1):
        if PLACEHOLDER.search(text):
            line_add(out, "阻断", f"{label} 第{i}文本段", "存在占位符或待办标记，成稿前须核实并处理。")
        if WORDY.search(text):
            line_add(out, "重大", f"{label} 第{i}文本段", "可能存在宣传性或绝对化用语，请依据证据修改。")
    return doc


def check_claims(doc, out):
    paras = contents(doc)
    claims = []
    for p in paras:
        mt = NUM.match(p)
        if mt:
            claims.append((int(mt.group(1)), mt.group(2)))
    if not claims:
        line_add(out, "阻断", "权利要求书", "未识别到阿拉伯数字顺序编号的权利要求（例如1.……）。")
        return
    indices = [v[0] for v in claims]
    if indices != list(range(1, len(indices) + 1)):
        line_add(out, "重大", "权利要求书", f"编号可能缺失/重复/不连续：{indices}。")
    multiple = set()
    for n, claim in claims:
        # 仅检查常见引用形式，复杂的宽范围引用必须人工核查。
        match = re.search(r"根据(.{0,55}?)所述", claim)
        if n > 1 and match:
            refregion = match.group(1)
            rnums = [int(v) for v in REF.findall(refregion)]
            if not rnums:
                rnums = [int(v) for v in re.findall(r"\d{1,3}", refregion)]
            if not rnums:
                line_add(out, "提示", f"权利要求{n}", "无法机械解析该从属引用，请人工核对。")
            if rnums and (0 in rnums or max(rnums) >= n):
                line_add(out, "重大", f"权利要求{n}", f"存在非在前引用：{rnums}。")
            if ('或' in refregion or '任一' in refregion or '至' in refregion or '-' in refregion) and len(rnums) >= 2:
                multiple.add(n)
        if n > 1 and not match:
            line_add(out, "提示", f"权利要求{n}", "无法识别引用形式；若该项为独立权利要求，请核查多个主题与单一性。")
        if re.search(r"(如说明书|如图\s*\d+|如附图)", claim):
            line_add(out, "重大", f"权利要求{n}", "有以说明书/附图代替技术特征的嫌疑。")
    if multiple:
        for n, claim in claims:
            if n in multiple:
                continue
            refs = [int(v) for v in re.findall(r"\d{1,3}", re.search(r"根据(.{0,55}?)所述", claim).group(1))] if re.search(r"根据(.{0,55}?)所述", claim) else []
            if any(x in multiple for x in refs) and ('或' in claim[:70] or '任一' in claim[:70]):
                line_add(out, "重大", f"权利要求{n}", "可能存在多项从属权利要求引用另一多项从属权利要求的情况。")
    line_add(out, "信息", "权利要求书", f"识别 {len(claims)} 项权利要求；自动检测无法判断必要技术特征、单一性及创造性。")


def check_spec(doc, out):
    raw = "\n".join(contents(doc))
    for h in HEADINGS:
        if h not in raw:
            if h == "附图说明":
                line_add(out, "提示", "说明书", "没有检测到附图说明；无附图的发明可能不需要该章节。")
            else:
                line_add(out, "重大", "说明书", f"未检测到章节“{h}”。")
    first = [(h, raw.find(h)) for h in HEADINGS if raw.find(h) >= 0]
    if [v[1] for v in first] != sorted(v[1] for v in first):
        line_add(out, "重大", "说明书", "章节标题顺序与常见五段式顺序不一致，请复核是否具有合法且必要的例外。")
    figureids = {int(k) for k in FIG.findall(raw)}
    if figureids and sorted(figureids) != list(range(1, max(figureids) + 1)):
        line_add(out, "提示", "说明书", f"正文识别图号 {sorted(figureids)}，存在跳号。")
    return figureids


def check_abstract(doc, out):
    paras = contents(doc)
    raw = "".join(p for p in paras if p.replace(' ', '') not in ('说明书摘要', '说 明 书 摘 要'))
    if raw.startswith('说明书摘要'):
        raw = raw[len('说明书摘要'):]
    if len(raw) > 300:
        line_add(out, "重大", "说明书摘要", f"摘要初步统计 {len(raw)} 字符（计标点），超过300。请以提交文本实际口径复核。")
    else:
        line_add(out, "信息", "说明书摘要", f"摘要初步统计 {len(raw)} 字符（计标点）。")
    if len(raw.strip()) < 30:
        line_add(out, "提示", "说明书摘要", "正文很短或缺失；确认是否明确核心技术方案。")


def write_report(out_path, issues, examined):
    level = {"阻断": 0, "重大": 1, "提示": 2, "信息": 3}
    issues.sort(key=lambda x: level.get(x[0], 5))
    counts = {x: sum(1 for i in issues if i[0] == x) for x in level}
    lines = ["# 发明专利文件自动初查报告", "", "> 自动规则不构成合规或授权法律结论；需人工确认技术真实性、权利要求支持、附图内容与渲染质量。", "", f"**已检查文件**：{', '.join(examined)}", "", f"**问题统计**：阻断 {counts['阻断']}，重大 {counts['重大']}，提示 {counts['提示']}，信息 {counts['信息']}。", "", "| 级别 | 位置 | 观察结果/操作建议 |", "|---|---|---|"]
    for severity, loc, msg in issues:
        lines.append(f"| {severity} | {loc.replace('|','/')} | {msg.replace('|','/').replace(chr(10),' ')} |")
    lines += ["", "## 本工具不能检查的事项", "", "- 新颖性、创造性、单一性、充分公开等实质法律问题；", "- 语义上的技术矛盾、技术效果真实性、引用公开文件的准确性；", "- 图片实际内容、线条可读性、页码下沿位置、批注与全部隐含格式；", "- 具体申请渠道的当前电子上传格式与官方表格要求。", "", "**定稿流程**：发明人确认事实 → 专业检索/代理核查 → 全文交叉审查 → 逐页图片渲染验收。"]
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return counts


def main():
    ap = argparse.ArgumentParser(description="中国发明专利 DOCX 基础规则审查（非法律结论）")
    ap.add_argument('--claims', type=Path, help='权利要求书 .docx')
    ap.add_argument('--spec', type=Path, help='说明书 .docx')
    ap.add_argument('--abstract', type=Path, help='说明书摘要 .docx')
    ap.add_argument('--out', type=Path, default=Path('专利自动初查报告.md'))
    args = ap.parse_args()
    if not any((args.claims, args.spec, args.abstract)):
        ap.error('至少提供 --claims / --spec / --abstract 之一')
    issues = []
    examined = []
    for title, path, fn in (("权利要求书", args.claims, check_claims), ("说明书", args.spec, check_spec), ("说明书摘要", args.abstract, check_abstract)):
        if path:
            if not path.is_file():
                ap.error(f'找不到 {path}')
            examined.append(f'{title} ({path.name})')
            doc = check_base(path, title, issues)
            fn(doc, issues)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    counts = write_report(args.out, issues, examined)
    print(f"已生成 {args.out} | 阻断{counts['阻断']} 重大{counts['重大']} 提示{counts['提示']}")

if __name__ == '__main__':
    main()
