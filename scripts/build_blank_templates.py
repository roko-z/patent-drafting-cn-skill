#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成不包含任何真实项目细节的专利 DOCX 工作模板。

用法: python build_blank_templates.py
依赖: python-docx
注意：生成的文件是工作底稿，不是官方请求书/最终可直接提交的文件。
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / 'templates' / 'docx'
OUT.mkdir(exist_ok=True, parents=True)


def apply_font(style, size=12, name='宋体'):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)


def base_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = Cm(2.5)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(1.5)
    s.bottom_margin = Cm(1.5)
    s.header_distance = Cm(1.15)
    s.footer_distance = Cm(1.2)
    apply_font(d.styles['Normal'])
    pf = d.styles['Normal'].paragraph_format
    pf.line_spacing = 1.5
    pf.space_before, pf.space_after = Pt(0), Pt(0)
    apply_font(d.styles['Title'], 12)
    apply_font(d.styles['Heading 1'], 12)
    d.styles['Heading 1'].font.color.rgb = RGBColor(0, 0, 0)
    # 各文件页码从1开始，默认不用自动首页混排
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.style = d.styles['Normal']
    run = footer.add_run()
    fld_begin = OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fld_sep = OxmlElement('w:fldChar'); fld_sep.set(qn('w:fldCharType'), 'separate')
    cached = OxmlElement('w:t'); cached.text = '1'
    fld_end = OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'), 'end')
    for item in (fld_begin, instr, fld_sep, cached, fld_end):run._r.append(item)
    return d


def line(d, content, heading=False, center=False):
    p = d.add_paragraph(style='Heading 1' if heading else 'Normal')
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if heading: p.paragraph_format.space_before = Pt(5)
    p.add_run(content)
    return p


def save(d, name):
    d.save(OUT / name)
    print('created', OUT / name)


def abstract_doc():
    d=base_doc(); line(d, '说明书摘要', center=True)
    line(d, '本发明公开了一种【技术客体名称】，属于【技术领域】。所述【技术客体】包括【必要技术构成】，所述【构件A】与【构件B】之间具有【明确的技术关系】，用于解决【待解决的技术问题】。通过【关键操作或结构作用】，获得【有事实依据的技术效果】，适用于【主要用途】。')
    save(d, '说明书摘要_空白模板.docx')


def claims_doc():
    d=base_doc(); line(d, '权利要求书', center=True)
    line(d, '1. 一种【技术客体名称】，包括【构件A】和【构件B】，其特征在于，所述【构件A】通过【具体连接关系】与所述【构件B】连接；所述【构件B】在【明确条件】下执行【技术操作】，得到【技术中间结果】。')
    line(d, '2. 根据权利要求1所述的【技术客体名称】，其特征在于，所述【构件A】包括【可选子构件】。')
    line(d, '3. 根据权利要求1所述的【技术客体名称】，其特征在于，所述【技术操作】包括【真实且已确认的进一步限定】。')
    save(d, '权利要求书_空白模板.docx')


def specification_doc():
    d=base_doc(); line(d, '说明书', center=True)
    line(d, '一种【技术客体名称】', center=True)
    line(d, '技术领域', heading=True)
    line(d, '本发明属于【技术领域】，具体涉及一种【技术客体名称】。')
    line(d, '背景技术', heading=True)
    line(d, '现有的【技术类型】通常采用【已核实的公开技术手段】。然而，在【具体应用条件】下仍存在【有证据的技术缺陷】。因此，有必要解决【进一步明确的技术问题】。')
    line(d, '发明内容', heading=True)
    line(d, '本发明所要解决的技术问题是【明确且可判断的问题】。')
    line(d, '为解决上述技术问题，提供一种【技术客体名称】，其包括【必要构件和具体关系】，并通过【具体手段】实现【中间技术结果】。')
    line(d, '优选地，【进一步限定特征；需要原始材料支持】。')
    line(d, '与现有技术相比，本发明在【相同可比的测试条件】下能够【由证据支持的客观技术效果】。')
    line(d, '附图说明', heading=True)
    line(d, '图1为本发明【整体结构/流程】的示意图。')
    line(d, '图2为本发明【关键局部关系】的示意图。')
    line(d, '图中：1—【构件A】；2—【构件B】；3—【构件C】。')
    line(d, '具体实施方式', heading=True)
    line(d, '下面结合附图与实际实施例，对本发明的技术方案作清楚、完整的说明。以下涉及的构件、参数、步骤、测试条件均应由发明人和原始技术资料核实。')
    line(d, '如图1所示，【构件A】1与【构件B】2通过【具体关系】连接。【信号/介质/动作】由【输入位置】进入，依次经过【必要操作】并在【输出位置】得到【结果】。')
    line(d, '在一个可实施例中：步骤S1，【输入条件和动作】；步骤S2，【依赖前一步输出的操作及参数】；步骤S3，【最终输出/控制过程】。')
    line(d, '在另一个可实施例中，【已证实的等效/优选方案】。其中【具体参数】的单位为【单位】，测试方法为【测试方法】，数据来源为【报告编号】。')
    save(d, '说明书_空白模板.docx')


def drawings_doc():
    d=base_doc(); line(d, '说明书附图（内部占位稿，请用正式图纸替换）',center=True)
    line(d, '图1', center=True)
    line(d, '【此处插入清晰、对应说明书的整体图，不应直接以此占位文档提交】', center=True)
    line(d, '图2',center=True)
    line(d, '【此处插入清晰、对应说明书的关键局部图，确认全部标号与说明书一致】',center=True)
    save(d, '说明书附图_占位模板.docx')


def disclosure_doc():
    d=base_doc(); line(d,'发明专利技术交底书（内部工作模板）',center=True)
    line(d,'本模板用于发明人和技术人员填写，暂不属于正式申请文件。未确认的数据请保留占位符，不得自行编造。')
    rows = [('拟定发明名称','【填写】'),('技术领域','【填写】'),('申请人/发明人/权属','【有权人员核实】'),('公开情况','【是否已有线上、线下公开，时间与渠道】'),('材料编号与版本','【填写】')]
    table=d.add_table(rows=0, cols=2); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for k,v in rows:
        cells=table.add_row().cells; cells[0].text=k; cells[1].text=v
    sections=[('一、现有技术与其不足','【最接近技术、公开文献、具体缺点】'),('二、拟解决技术问题','【具体客观的技术问题】'),('三、完整技术方案','【必要部件/步骤、作用关系、操作原理】'),('四、实施条件与参数','【已验证的范围、单位、条件、检测/测试方法】'),('五、技术效果和对比例','【真实数据、基线、实验条件与证据位置】'),('六、优选与替代方案','【已验证的替代实现，不得臆测】'),('七、附图与标号清单','【图1、图2……及部件编号表】'),('八、检索与区别特征','【D1、D2、专利公开号及区别】'),('九、待发明人确认','【P0/P1/P2 待确认清单及责任人】')]
    for h,t in sections: line(d,h,heading=True);line(d,t)
    save(d, '技术交底书_空白模板.docx')

if __name__ == '__main__':
    for f in (abstract_doc,claims_doc,specification_doc,drawings_doc,disclosure_doc):f()
